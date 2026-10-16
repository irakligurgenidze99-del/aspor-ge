#!/usr/bin/env python3
"""
One-off data preparation: converts the supplier workbook "Aspor (2).xlsx" (sheet 2)
into seed/aspor-catalog.json + seed/images/*.webp for the DEMO seed.

Deliberately NOT read: VIP/purchase price, cost price, B2B/B2C markup & margin,
competitor columns, quantities and carton/logistics data. Those are internal
and must never reach the public site. Wholesale prices are left empty.

Usage: python3 scripts/extract_aspor_xlsx.py "/path/to/Aspor (2).xlsx"
"""
import io, json, re, sys
from collections import defaultdict
import openpyxl
from PIL import Image

SRC = sys.argv[1]
OUT = 'seed'
wb = openpyxl.load_workbook(SRC)
ws = wb.worksheets[1]  # "Sheet1 (2)" — has headers MODEL NO., PICTURE, PACKAGE, SPECIFICATIONS, COLOR

imgs = defaultdict(dict)
for im in ws._images:
    a = im.anchor._from
    imgs[a.row + 1][a.col] = im

COMMERCIAL_NOTES = {'no extra discount', 'limited stock'}
COLOR_KA = {'black': 'შავი', 'white': 'თეთრი', 'gray': 'ნაცრისფერი', 'grey': 'ნაცრისფერი', 'gry': 'ნაცრისფერი',
            'dark gray': 'მუქი ნაცრისფერი', 'blue': 'ლურჯი', 'purple': 'იისფერი'}
COLOR_EN = {'gry': 'gray', 'grey': 'gray'}

def slug(s):
    return re.sub(r'[^A-Za-z0-9]+', '-', s).strip('-').upper()

def category(name, specs, base):
    t = (name + ' ' + specs).lower()
    if base.startswith(('A78',)) or 'hub' in t or base == 'A783':
        return 'other'
    if 'mp3' in t or 'power socket' in t or base in ('A502', 'A503', 'A505 PD'):
        return 'other'
    if 'power bank' in t or 'powerbank' in t or 'magsafe' in t or re.search(r'capa[ct]i[ct]y\s*[:：]?\s*\d+\s*mah', t):
        return 'power-banks'
    if base.startswith('AC') or ('cable' in name.lower() and 'charger' not in name.lower()) or (re.search(r'\bto (micro|iphone|lightning|type-c)\b', t) and 'length' in t and 'charger' not in t):
        return 'cables'
    return 'chargers'

def save_img(im, name):
    data = im._data()
    pil = Image.open(io.BytesIO(data)).convert('RGBA')
    bg = Image.new('RGBA', pil.size, (255, 255, 255, 255))
    bg.alpha_composite(pil)
    out = bg.convert('RGB')
    out.thumbnail((1400, 1400))
    path = f'{OUT}/images/{name}.webp'
    out.save(path, 'WEBP', quality=82, method=6)
    return path

products, seen = [], set()
for r in range(5, ws.max_row + 1):
    model = ws.cell(r, 1).value
    specs = (ws.cell(r, 4).value or '').strip()
    colors = (ws.cell(r, 5).value or '').strip()
    weight = ws.cell(r, 14).value
    if not model:
        continue
    lines = [l.strip() for l in str(model).split('\n') if l.strip()]
    base = lines[0]
    notes = [l for l in lines[1:] if l.lower() not in COMMERCIAL_NOTES]
    cert = [l for l in notes if l.upper() == 'CCC']
    length = [l for l in notes if re.fullmatch(r'\d+M', l, re.I)]
    power_note = [l for l in notes if re.match(r'^\d+(\.\d+)?W\b', l)]
    variant = [l for l in notes if l not in cert and l not in length and l not in power_note]
    spec_lines = [l.strip() for l in specs.split('\n') if l.strip()]
    first = spec_lines[0] if spec_lines else ''
    name_en = first if first and ':' not in first and '：' not in first else ''
    rest = spec_lines[1:] if name_en else spec_lines
    if not name_en:
        cand = next((l for l in spec_lines if ':' not in l and '：' not in l and re.search(r'[A-Za-z]{4,}', l)), '')
        if cand:
            name_en = cand
            rest = [l for l in spec_lines if l != cand]
    attrs, feats = [], []
    for l in rest:
        m = re.match(r'^\s*([^:：]{1,60})[:：]\s*(.+)$', l)
        if m:
            attrs.append({'label': m.group(1).strip(), 'value': m.group(2).strip()})
        else:
            feats.append(l)
    if variant:
        attrs.insert(0, {'label': 'კომპლექტაცია', 'value': ' / '.join(variant)})
    if length:
        attrs.append({'label': 'სიგრძე', 'value': length[0].lower()})
    if cert:
        attrs.append({'label': 'სერტიფიკაცია', 'value': 'CCC'})
    for f in feats:
        attrs.append({'label': 'ფუნქცია', 'value': f})
    if isinstance(weight, str) and re.search(r'\d\s*g$', weight.strip()):
        attrs.append({'label': 'წონა', 'value': re.sub(r'\s+', ' ', weight.strip())})
    # de-duplicate labels
    cnt = defaultdict(int)
    for a in attrs:
        cnt[a['label']] += 1
        if cnt[a['label']] > 1:
            a['label'] = f"{a['label']} {cnt[a['label']]}"

    cat = category(name_en, specs, base)
    color_list = [c.strip() for c in re.split(r'[\n/]', colors) if c.strip()] or ['']
    pics = []
    if 1 in imgs[r]:
        pics.append(save_img(imgs[r][1], f'r{r}-product'))
    if 2 in imgs[r]:
        pics.append(save_img(imgs[r][2], f'r{r}-package'))

    model_number = base + (f" ({' / '.join(variant)})" if variant else '')
    sku_base = slug(base) + ('-' + slug(' '.join(variant)) if variant else '')
    for c in color_list:
        ce = COLOR_EN.get(c.lower(), c.lower())
        sku = sku_base + ('-' + slug(ce) if ce else '')
        if sku in seen:
            sku += f'-R{r}'
        seen.add(sku)
        ka_color = COLOR_KA.get(c.lower(), ce)
        products.append({
            'sku': sku, 'row': r, 'category': cat, 'modelNumber': model_number,
            'nameEn': name_en, 'nameKa': f"ASPOR {base} {name_en}".strip(), 'color': ka_color, 'colorEn': ce,
            'specs': attrs, 'images': pics,
        })

json.dump({'source': 'Aspor (2).xlsx — sheet "Sheet1 (2)"', 'note': 'DEMO seed. SKUs with colour suffix are PROVISIONAL and must be confirmed. No prices imported.', 'products': products}, open(f'{OUT}/aspor-catalog.json', 'w'), ensure_ascii=False, indent=1)
print(len(products), 'products', sum(1 for p in products if p['category'] == 'other'), 'other')
